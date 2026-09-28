from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
import sqlite3, os, io, math, re
from datetime import datetime, date
from functools import wraps

app = Flask(__name__)
app.secret_key = "warkop-sederhana-secret-2026"
BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(BASE, "warkop.db")

MENUS = [
    ("Kopi Hitam", 5000, "gelas", "Kopi bubuk, gula pasir, air mineral", "coffee"),
    ("Nasi Goreng", 15000, "porsi", "Beras, telur, bawang, kecap, minyak, bumbu", "nasi"),
    ("Teh Manis", 4000, "gelas", "Teh celup, gula pasir, air mineral", "tea"),
    ("Good Day", 5000, "gelas", "Good Day sachet, air mineral", "goodday"),
    ("Pisang Goreng", 8000, "porsi", "Pisang, tepung, gula, minyak", "banana"),
    ("Mie Goreng", 10000, "porsi", "Mie instan, telur, minyak, bumbu mie", "noodle"),
    ("Es Jeruk", 6000, "gelas", "Jeruk, gula pasir, air mineral", "orange"),
    ("Air Mineral", 4000, "botol", "Air mineral botol", "water"),
]

RECIPES = {
    "Kopi Hitam": [("Kopi bubuk", "gram", 12), ("Gula pasir", "gram", 10), ("Air mineral", "ml", 180)],
    "Nasi Goreng": [("Beras", "gram", 100), ("Telur", "butir", 1), ("Bawang putih", "gram", 5), ("Bawang merah", "gram", 8), ("Kecap manis", "ml", 10), ("Minyak goreng", "ml", 15), ("Bumbu nasi goreng", "gram", 10)],
    "Teh Manis": [("Teh celup", "sachet", 1), ("Gula pasir", "gram", 15), ("Air mineral", "ml", 200)],
    "Good Day": [("Good Day sachet", "sachet", 1), ("Air mineral", "ml", 180)],
    "Pisang Goreng": [("Pisang", "buah", 3), ("Tepung terigu", "gram", 50), ("Gula pasir", "gram", 5), ("Minyak goreng", "ml", 40)],
    "Mie Goreng": [("Mie instan", "bungkus", 1), ("Telur", "butir", 1), ("Minyak goreng", "ml", 10), ("Bumbu mie", "gram", 5)],
    "Es Jeruk": [("Jeruk peras", "buah", 1), ("Gula pasir", "gram", 15), ("Air mineral", "ml", 180)],
    "Air Mineral": [("Air mineral botol", "botol", 1)],
}

STOCK = [
    ("Kopi bubuk", "gram", 1200, 500), ("Gula pasir", "gram", 3500, 700), ("Air mineral", "ml", 30000, 5000),
    ("Beras", "gram", 9000, 2500), ("Telur", "butir", 35, 10), ("Bawang putih", "gram", 900, 150),
    ("Bawang merah", "gram", 1200, 200), ("Kecap manis", "ml", 3000, 500), ("Minyak goreng", "ml", 1200, 1500),
    ("Bumbu nasi goreng", "gram", 500, 150), ("Teh celup", "sachet", 80, 20), ("Good Day sachet", "sachet", 18, 20),
    ("Pisang", "buah", 18, 15), ("Tepung terigu", "gram", 4000, 700), ("Mie instan", "bungkus", 40, 10),
    ("Bumbu mie", "gram", 300, 50), ("Jeruk peras", "buah", 28, 8), ("Air mineral botol", "botol", 35, 8)
]


def connect():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c


def ensure_column(c, table, column, definition):
    cols = [r[1] for r in c.execute(f"PRAGMA table_info({table})").fetchall()]
    if column not in cols:
        c.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_db():
    c = connect()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE, password TEXT);
    CREATE TABLE IF NOT EXISTS menus(id INTEGER PRIMARY KEY, name TEXT UNIQUE, price INTEGER, unit TEXT, description TEXT, image TEXT);
    CREATE TABLE IF NOT EXISTS stock(id INTEGER PRIMARY KEY, name TEXT UNIQUE, unit TEXT, qty REAL, min_qty REAL);
    CREATE TABLE IF NOT EXISTS recipes(id INTEGER PRIMARY KEY, menu_name TEXT, material TEXT, unit TEXT, qty REAL);
    CREATE TABLE IF NOT EXISTS sales(id INTEGER PRIMARY KEY, invoice TEXT, sale_date TEXT, total INTEGER);
    CREATE TABLE IF NOT EXISTS sale_items(id INTEGER PRIMARY KEY, sale_id INTEGER, menu_name TEXT, qty INTEGER, price INTEGER);
    CREATE TABLE IF NOT EXISTS stock_logs(id INTEGER PRIMARY KEY, material TEXT, type TEXT, qty REAL, unit TEXT, note TEXT, log_date TEXT, menu_name TEXT);
    CREATE TABLE IF NOT EXISTS forecast_runs(id INTEGER PRIMARY KEY, start_date TEXT, end_date TEXT, horizon INTEGER, method TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS forecast_items(id INTEGER PRIMARY KEY, run_id INTEGER, material TEXT, unit TEXT, current REAL, min_qty REAL, predicted REAL, recommended REAL, status TEXT);
    """)
    # Migrate databases created by older versions.
    ensure_column(c, "stock_logs", "unit", "TEXT")
    ensure_column(c, "stock_logs", "menu_name", "TEXT")
    # Fill missing units in old stock logs from current stock table.
    c.execute("UPDATE stock_logs SET unit=(SELECT unit FROM stock WHERE stock.name=stock_logs.material) WHERE unit IS NULL OR unit=''" )
    c.execute("INSERT OR IGNORE INTO users(username,password) VALUES('admin','admin123')")
    for name, price, unit, desc, img in MENUS:
        c.execute("INSERT OR IGNORE INTO menus(name,price,unit,description,image) VALUES(?,?,?,?,?)", (name, price, unit, desc, img))
    for name, unit, qty, min_qty in STOCK:
        c.execute("INSERT OR IGNORE INTO stock(name,unit,qty,min_qty) VALUES(?,?,?,?)", (name, unit, qty, min_qty))
    for menu, items in RECIPES.items():
        for material, unit, qty in items:
            row = c.execute("SELECT id FROM recipes WHERE menu_name=? AND material=?", (menu, material)).fetchone()
            if not row:
                c.execute("INSERT INTO recipes(menu_name,material,unit,qty) VALUES(?,?,?,?)", (menu, material, unit, qty))
    # Backfill old log menu names from the note where possible.
    old_logs = c.execute("SELECT id,note FROM stock_logs WHERE (menu_name IS NULL OR menu_name='') AND note LIKE 'Penjualan %'").fetchall()
    for r in old_logs:
        m = re.match(r"Penjualan (.+?) x[0-9.]+$", r["note"] or "")
        if m:
            c.execute("UPDATE stock_logs SET menu_name=? WHERE id=?", (m.group(1), r["id"]))
    c.commit()
    c.close()


def auth_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user"):
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapper


def parse_date_range(default_month=True):
    today = date.today()
    default_start = today.replace(day=1).isoformat() if default_month else ""
    default_end = today.isoformat() if default_month else ""
    start = request.args.get("start") or default_start
    end = request.args.get("end") or default_end
    return start, end


def validate_range(start, end):
    try:
        d1 = datetime.strptime(start, "%Y-%m-%d").date()
        d2 = datetime.strptime(end, "%Y-%m-%d").date()
        return d1, d2
    except ValueError:
        return None, None


def date_range_days(start, end):
    d1, d2 = validate_range(start, end)
    return max(1, (d2 - d1).days + 1) if d1 and d2 else 1


def money(n):
    return "Rp {:,.0f}".format(n).replace(",", ".")


def fmt_qty(n):
    n = float(n or 0)
    return str(int(n)) if n.is_integer() else f"{n:.1f}"


def get_sales_rows(start, end):
    c = connect()
    rows = c.execute("""SELECT s.invoice,s.sale_date,si.menu_name,si.qty,si.price,(si.qty*si.price) line_total
                        FROM sale_items si JOIN sales s ON s.id=si.sale_id
                        WHERE date(s.sale_date) BETWEEN date(?) AND date(?)
                        ORDER BY s.sale_date DESC,s.id DESC,si.id""", (start, end)).fetchall()
    c.close()
    return rows


def get_bahan_detail(start, end):
    c = connect()
    rows = c.execute("""SELECT date(log_date) log_day, COALESCE(menu_name,'-') menu_name, material,
                               COALESCE(unit,(SELECT unit FROM stock s WHERE s.name=stock_logs.material),'') unit,
                               SUM(qty) qty
                        FROM stock_logs
                        WHERE type='KELUAR' AND date(log_date) BETWEEN date(?) AND date(?)
                        GROUP BY date(log_date), menu_name, material, unit
                        ORDER BY log_day DESC, menu_name, material""", (start, end)).fetchall()
    c.close()
    return rows


def get_bahan_rekap_menu(start, end):
    c = connect()
    rows = c.execute("""SELECT COALESCE(menu_name,'-') menu_name, material,
                               COALESCE(unit,(SELECT unit FROM stock s WHERE s.name=stock_logs.material),'') unit,
                               SUM(qty) qty
                        FROM stock_logs
                        WHERE type='KELUAR' AND date(log_date) BETWEEN date(?) AND date(?)
                        GROUP BY menu_name, material, unit
                        ORDER BY menu_name, material""", (start, end)).fetchall()
    c.close()
    return rows


def get_bahan_rekap_material(start, end):
    c = connect()
    rows = c.execute("""SELECT material,
                               COALESCE(unit,(SELECT unit FROM stock s WHERE s.name=stock_logs.material),'') unit,
                               SUM(qty) qty
                        FROM stock_logs
                        WHERE type='KELUAR' AND date(log_date) BETWEEN date(?) AND date(?)
                        GROUP BY material, unit
                        ORDER BY material""", (start, end)).fetchall()
    c.close()
    return rows


def get_grouped_recipes():
    c = connect()
    rows = c.execute("""SELECT r.menu_name,r.material,r.unit,r.qty,m.price,m.unit menu_unit
                        FROM recipes r LEFT JOIN menus m ON m.name=r.menu_name
                        ORDER BY m.id,r.id""").fetchall()
    c.close()
    groups = []
    by_menu = {}
    for r in rows:
        if r["menu_name"] not in by_menu:
            group = {"menu_name": r["menu_name"], "menu_unit": r["menu_unit"], "price": r["price"], "items": []}
            by_menu[r["menu_name"]] = group
            groups.append(group)
        by_menu[r["menu_name"]]["items"].append(r)
    return groups


def daily_menu_sales(start, end):
    c = connect()
    rows = c.execute("""SELECT date(s.sale_date) sale_day, si.menu_name, SUM(si.qty) qty
                      FROM sale_items si JOIN sales s ON s.id=si.sale_id
                      WHERE date(s.sale_date) BETWEEN date(?) AND date(?)
                      GROUP BY date(s.sale_date),si.menu_name ORDER BY sale_day""", (start, end)).fetchall()
    c.close()
    return rows


def forecast_materials(start, end, horizon):
    d1, d2 = validate_range(start, end)
    if not d1 or not d2 or d2 < d1:
        raise ValueError("Periode prediksi tidak valid.")
    c = connect()
    stocks = c.execute("SELECT * FROM stock ORDER BY name").fetchall()
    all_menu_names = [r["name"] for r in c.execute("SELECT name FROM menus ORDER BY id").fetchall()]
    menu_rows = c.execute("""SELECT si.menu_name,SUM(si.qty) total
                             FROM sale_items si JOIN sales s ON s.id=si.sale_id
                             WHERE date(s.sale_date) BETWEEN date(?) AND date(?)
                             GROUP BY si.menu_name""", (start, end)).fetchall()
    c.close()
    days = max(1, (d2 - d1).days + 1)
    menu_totals = {r["menu_name"]: float(r["total"]) for r in menu_rows}
    daily = daily_menu_sales(start, end)
    menu_pred = {}
    used_prophet = False
    try:
        from prophet import Prophet
        import pandas as pd
        idx = pd.date_range(start=start, end=end, freq="D")
        for menu in all_menu_names:
            values = {str(r["sale_day"]): float(r["qty"]) for r in daily if r["menu_name"] == menu}
            df = pd.DataFrame({"ds": idx, "y": [values.get(d.strftime("%Y-%m-%d"), 0.0) for d in idx]})
            if df["y"].sum() <= 0 or len(df) < 2:
                menu_pred[menu] = menu_totals.get(menu, 0.0) / days * horizon
                continue
            model = Prophet(daily_seasonality=False, weekly_seasonality=True, yearly_seasonality=False)
            model.fit(df)
            future = model.make_future_dataframe(periods=horizon, include_history=False)
            pred = model.predict(future)["yhat"].clip(lower=0).sum()
            menu_pred[menu] = float(pred)
            used_prophet = True
    except Exception:
        for menu in all_menu_names:
            menu_pred[menu] = menu_totals.get(menu, 0.0) / days * horizon
    output = []
    for s in stocks:
        predicted = 0.0
        for menu, pred in menu_pred.items():
            for material, unit, qty in RECIPES.get(menu, []):
                if material == s["name"]:
                    predicted += pred * qty
        recommended = max(0, math.ceil(predicted + s["min_qty"] - s["qty"]))
        output.append({
            "material": s["name"], "unit": s["unit"], "current": s["qty"], "min": s["min_qty"],
            "predicted": round(predicted, 1), "remaining": round(s["qty"] - predicted, 1), "recommended": recommended,
            "status": "PERLU RESTOCK" if recommended > 0 else "AMAN"
        })
    return output, used_prophet


def save_forecast(start, end, horizon, method, rows):
    c = connect()
    run_id = c.execute("INSERT INTO forecast_runs(start_date,end_date,horizon,method,created_at) VALUES(?,?,?,?,?)",
                       (start, end, horizon, method, datetime.now().isoformat(timespec="seconds"))).lastrowid
    for r in rows:
        c.execute("""INSERT INTO forecast_items(run_id,material,unit,current,min_qty,predicted,recommended,status)
                     VALUES(?,?,?,?,?,?,?,?)""", (run_id, r["material"], r["unit"], r["current"], r["min"], r["predicted"], r["recommended"], r["status"]))
    c.commit(); c.close()
    return run_id


@app.route("/", methods=["GET", "POST"])
def login():
    if session.get("user"):
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        c = connect()
        row = c.execute("SELECT * FROM users WHERE username=? AND password=?", (request.form["username"], request.form["password"])).fetchone()
        c.close()
        if row:
            session["user"] = row["username"]
            return redirect(url_for("dashboard"))
        flash("Username atau password salah.")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear(); return redirect(url_for("login"))


@app.route("/dashboard")
@auth_required
def dashboard():
    c = connect()
    today = date.today()
    # Default dashboard period = current month. User can change the period.
    default_start = today.replace(day=1).isoformat()
    default_end = today.isoformat()
    start_date = request.args.get("start_date") or default_start
    end_date = request.args.get("end_date") or default_end
    try:
        datetime.strptime(start_date, "%Y-%m-%d")
        datetime.strptime(end_date, "%Y-%m-%d")
        if start_date > end_date:
            raise ValueError
    except ValueError:
        start_date, end_date = default_start, default_end
        flash("Periode dashboard tidak valid. Dikembalikan ke periode bulan berjalan.")

    revenue = c.execute(
        "SELECT COALESCE(SUM(total),0) x FROM sales WHERE date(sale_date) BETWEEN date(?) AND date(?)",
        (start_date, end_date)
    ).fetchone()["x"]
    tx = c.execute(
        "SELECT COUNT(*) x FROM sales WHERE date(sale_date) BETWEEN date(?) AND date(?)",
        (start_date, end_date)
    ).fetchone()["x"]
    menu_qty = c.execute(
        "SELECT COALESCE(SUM(si.qty),0) x FROM sale_items si JOIN sales s ON s.id=si.sale_id "
        "WHERE date(s.sale_date) BETWEEN date(?) AND date(?)",
        (start_date, end_date)
    ).fetchone()["x"]
    low = c.execute("SELECT COUNT(*) x FROM stock WHERE qty<=min_qty").fetchone()["x"]
    safe = c.execute("SELECT COUNT(*) x FROM stock WHERE qty>min_qty").fetchone()["x"]
    low_rows = c.execute("SELECT * FROM stock WHERE qty<=min_qty ORDER BY qty ASC").fetchall()
    menus = c.execute("SELECT * FROM menus ORDER BY id").fetchall()
    c.close()
    return render_template(
        "dashboard.html", title="Dashboard", revenue=revenue, tx=tx, menu_qty=menu_qty,
        low=low, safe=safe, low_rows=low_rows, menus=menus,
        start_date=start_date, end_date=end_date
    )


@app.route("/kasir", methods=["GET", "POST"])
@auth_required
def kasir():
    c = connect(); menus = c.execute("SELECT * FROM menus ORDER BY id").fetchall()
    if request.method == "POST":
        cart = []
        for m in menus:
            qty = int(request.form.get("qty_" + str(m["id"]), 0) or 0)
            if qty > 0:
                cart.append((m["name"], qty, m["price"]))
        sale_date = request.form.get("sale_date") or date.today().isoformat()
        if not cart:
            c.close(); flash("Pilih minimal satu menu."); return redirect(url_for("kasir"))
        try:
            datetime.strptime(sale_date, "%Y-%m-%d")
        except ValueError:
            c.close(); flash("Tanggal transaksi tidak valid."); return redirect(url_for("kasir"))
        needed = {}
        for name, qty, price in cart:
            for material, unit, per_item in RECIPES[name]:
                needed[material] = needed.get(material, 0) + qty * per_item
        for material, need in needed.items():
            row = c.execute("SELECT qty,unit FROM stock WHERE name=?", (material,)).fetchone()
            if not row or row["qty"] < need:
                c.close(); flash(f"Stok {material} tidak cukup. Tersedia {fmt_qty(row['qty']) if row else '0'} {row['unit'] if row else ''}, dibutuhkan {fmt_qty(need)}."); return redirect(url_for("kasir"))
        total = sum(q * p for _, q, p in cart)
        invoice = "TRX-" + datetime.now().strftime("%Y%m%d%H%M%S%f")
        sid = c.execute("INSERT INTO sales(invoice,sale_date,total) VALUES(?,?,?)", (invoice, sale_date, total)).lastrowid
        for name, qty, price in cart:
            c.execute("INSERT INTO sale_items(sale_id,menu_name,qty,price) VALUES(?,?,?,?)", (sid, name, qty, price))
            for material, unit, per_item in RECIPES[name]:
                used = qty * per_item
                c.execute("UPDATE stock SET qty=qty-? WHERE name=?", (used, material))
                c.execute("INSERT INTO stock_logs(material,type,qty,unit,note,log_date,menu_name) VALUES(?,?,?,?,?,?,?)",
                          (material, "KELUAR", used, unit, f"Penjualan {name} x{qty}", sale_date + " 12:00:00", name))
        c.commit(); c.close()
        flash(f"Transaksi {invoice} tersimpan. Grand Total {money(total)}.")
        return redirect(url_for("kasir"))
    c.close()
    return render_template("kasir.html", title="Kasir", menus=menus, today=date.today().isoformat())


@app.route("/master-menu")
@auth_required
def master_menu():
    c = connect(); menus = c.execute("SELECT * FROM menus ORDER BY id").fetchall(); c.close()
    return render_template("master_menu.html", title="Master Menu", menus=menus)


@app.route("/penjualan")
@auth_required
def penjualan():
    start, end = parse_date_range(True); rows = get_sales_rows(start, end)
    return render_template("penjualan.html", title="Riwayat Penjualan", rows=rows, start=start, end=end)


@app.route("/stok")
@auth_required
def stok():
    start, end = parse_date_range(True)
    c = connect()
    rows = c.execute("SELECT * FROM stock ORDER BY name").fetchall()
    movements = c.execute("""SELECT date(log_date) log_day, material, type, qty, unit, note, COALESCE(menu_name,'-') menu_name
                             FROM stock_logs
                             WHERE date(log_date) BETWEEN date(?) AND date(?)
                             ORDER BY log_day DESC, id DESC""", (start, end)).fetchall()
    c.close()
    return render_template("stok.html", title="Stok Bahan", rows=rows, movements=movements, start=start, end=end, today=date.today().isoformat())


@app.route("/stok/tambah", methods=["POST"])
@auth_required
def tambah_stok():
    material = request.form["material"]; qty = float(request.form["qty"]); note = request.form.get("note") or "Stok masuk"; log_date = request.form.get("log_date") or date.today().isoformat()
    c = connect(); row = c.execute("SELECT unit FROM stock WHERE name=?", (material,)).fetchone()
    unit = row["unit"] if row else ""
    c.execute("UPDATE stock SET qty=qty+? WHERE name=?", (qty, material))
    c.execute("INSERT INTO stock_logs(material,type,qty,unit,note,log_date,menu_name) VALUES(?,?,?,?,?,?,NULL)", (material, "MASUK", qty, unit, note, log_date + " 12:00:00"))
    c.commit(); c.close(); flash("Stok berhasil ditambahkan."); return redirect(url_for("stok"))


@app.route("/bahan-terpakai")
@auth_required
def bahan_terpakai():
    start, end = parse_date_range(True)
    menu_rows = get_bahan_rekap_menu(start, end)
    material_rows = get_bahan_rekap_material(start, end)
    detail_rows = get_bahan_detail(start, end)
    return render_template("bahan.html", title="Bahan Terpakai", start=start, end=end, menu_rows=menu_rows, material_rows=material_rows, detail_rows=detail_rows)


@app.route("/resep")
@auth_required
def resep():
    groups = get_grouped_recipes()
    return render_template("resep.html", title="Resep Menu", groups=groups)


@app.route("/prediksi", methods=["GET", "POST"])
@auth_required
def prediksi():
    result = []; start = ""; end = ""; horizon = 7; method = ""; run_id = None
    if request.method == "POST":
        start = request.form["start"]; end = request.form["end"]; horizon = int(request.form["horizon"])
        d1, d2 = validate_range(start, end)
        if not d1 or not d2 or d2 < d1:
            flash("Periode tanggal prediksi tidak valid."); return redirect(url_for("prediksi"))
        result, used = forecast_materials(start, end, horizon)
        method = "Prophet" if used else "Rata-rata historis"
        run_id = save_forecast(start, end, horizon, method, result)
        session["last_forecast"] = {"start": start, "end": end, "horizon": horizon, "rows": result, "method": method, "run_id": run_id}
    else:
        last = session.get("last_forecast")
        if last:
            start, end, horizon, result, method, run_id = last["start"], last["end"], last["horizon"], last["rows"], last["method"], last.get("run_id")
        else:
            start, end = parse_date_range(True)
    c = connect(); history = c.execute("SELECT * FROM forecast_runs ORDER BY id DESC LIMIT 20").fetchall(); c.close()
    return render_template("prediksi.html", title="Prediksi Prophet", result=result, start=start, end=end, horizon=horizon, method=method, history=history, run_id=run_id)


@app.route("/prediksi/terapkan", methods=["POST"])
@auth_required
def terapkan():
    c = connect()
    for key, value in request.form.items():
        if key.startswith("rec_"):
            qty = float(value or 0); material = key[4:]
            if qty > 0:
                row = c.execute("SELECT unit FROM stock WHERE name=?", (material,)).fetchone(); unit = row["unit"] if row else ""
                c.execute("UPDATE stock SET qty=qty+? WHERE name=?", (qty, material))
                c.execute("INSERT INTO stock_logs(material,type,qty,unit,note,log_date,menu_name) VALUES(?,?,?,?,?,?,NULL)", (material, "MASUK", qty, unit, "Restock dari rekomendasi Prophet", datetime.now().isoformat(timespec="seconds")))
    c.commit(); c.close(); flash("Rekomendasi yang diisi sudah masuk ke stok."); return redirect(url_for("stok"))


@app.route("/reset-data", methods=["GET", "POST"])
@auth_required
def reset_data():
    if request.method == "GET":
        return render_template("reset.html", title="Reset Data Simulasi")
    if request.form.get("confirmation", "").strip().upper() != "RESET":
        flash("Reset dibatalkan. Ketik RESET untuk mengonfirmasi.")
        return redirect(url_for("dashboard"))
    c = connect()
    c.execute("DELETE FROM sale_items")
    c.execute("DELETE FROM sales")
    c.execute("DELETE FROM stock_logs")
    c.execute("DELETE FROM forecast_items")
    c.execute("DELETE FROM forecast_runs")
    for name, unit, qty, min_qty in STOCK:
        c.execute("UPDATE stock SET qty=?, unit=?, min_qty=? WHERE name=?", (qty, unit, min_qty, name))
    c.commit()
    c.close()
    session.pop("last_forecast", None)
    flash("Data simulasi berhasil direset. Riwayat kembali 0 dan stok dikembalikan ke stok awal.")
    return redirect(url_for("dashboard"))


# ---------- Exports ----------
def make_xlsx(filename, headers, rows, title="Laporan", subtitle=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = title[:31]
    ws.append([title]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers)); ws["A1"].font = Font(bold=True, size=14); ws["A1"].alignment = Alignment(horizontal="center")
    if subtitle:
        ws.append([subtitle]); ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers)); ws["A2"].alignment = Alignment(horizontal="center")
    header_row = 3 if subtitle else 2
    ws.append(headers)
    fill = PatternFill("solid", fgColor="754622"); white = Font(color="FFFFFF", bold=True)
    for cell in ws[header_row]: cell.fill = fill; cell.font = white; cell.alignment = Alignment(horizontal="center")
    for row in rows: ws.append(list(row))
    for col in range(1, len(headers)+1):
        letter = get_column_letter(col); max_len = max([len(str(ws.cell(r,col).value or "")) for r in range(1, ws.max_row+1)] + [10]); ws.column_dimensions[letter].width = min(max_len+3, 35)
    thin = Side(style="thin", color="DDDDDD")
    for row in ws.iter_rows(min_row=header_row):
        for cell in row: cell.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=header_row+1):
        for cell in row: cell.border = Border(bottom=thin); cell.alignment = Alignment(vertical="top")
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def make_pdf(filename, title, headers, rows, subtitle=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    bio = io.BytesIO(); doc = SimpleDocTemplate(bio, pagesize=landscape(A4), rightMargin=25, leftMargin=25, topMargin=25, bottomMargin=25)
    styles = getSampleStyleSheet(); styles["Title"].alignment = TA_CENTER
    story = [Paragraph(title, styles["Title"])]
    if subtitle: story += [Paragraph(subtitle, styles["Normal"]), Spacer(1, 10)]
    data = [headers] + [[str(v) for v in r] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#754622")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#DDDDDD")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("FONTSIZE", (0,0), (-1,-1), 8), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#FAF7F3")])]))
    story.append(table); doc.build(story); bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=filename, mimetype="application/pdf")


@app.route("/export/penjualan.xlsx")
@auth_required
def export_penjualan_xlsx():
    start, end = parse_date_range(True); rows = get_sales_rows(start, end)
    data = [(r["invoice"], r["sale_date"], r["menu_name"], r["qty"], money(r["price"]), money(r["line_total"])) for r in rows]
    return make_xlsx("riwayat_penjualan.xlsx", ["Invoice","Tanggal","Menu","Qty","Harga Satuan","Total"], data, "Riwayat Penjualan", f"Periode: {start} s.d. {end}")


@app.route("/export/penjualan.pdf")
@auth_required
def export_penjualan_pdf():
    start, end = parse_date_range(True); rows = get_sales_rows(start, end)
    data = [(r["invoice"], r["sale_date"], r["menu_name"], r["qty"], money(r["price"]), money(r["line_total"])) for r in rows]
    return make_pdf("riwayat_penjualan.pdf", "Riwayat Penjualan", ["Invoice","Tanggal","Menu","Qty","Harga Satuan","Total"], data, f"Periode: {start} s.d. {end}")


@app.route("/export/bahan.xlsx")
@auth_required
def export_bahan_xlsx():
    start, end = parse_date_range(True); rows = get_bahan_rekap_menu(start, end)
    data = [(r["menu_name"], r["material"], fmt_qty(r["qty"]), r["unit"]) for r in rows]
    return make_xlsx("riwayat_bahan_terpakai.xlsx", ["Menu","Bahan","Total Terpakai","Satuan"], data, "Riwayat Bahan Terpakai", f"Periode: {start} s.d. {end}")


@app.route("/export/bahan.pdf")
@auth_required
def export_bahan_pdf():
    start, end = parse_date_range(True); rows = get_bahan_rekap_menu(start, end)
    data = [(r["menu_name"], r["material"], fmt_qty(r["qty"]), r["unit"]) for r in rows]
    return make_pdf("riwayat_bahan_terpakai.pdf", "Riwayat Bahan Terpakai", ["Menu","Bahan","Total Terpakai","Satuan"], data, f"Periode: {start} s.d. {end}")


@app.route("/export/prediksi.xlsx")
@auth_required
def export_prediksi_xlsx():
    last = session.get("last_forecast")
    if not last: flash("Jalankan prediksi terlebih dahulu."); return redirect(url_for("prediksi"))
    data = [(r["material"], fmt_qty(r["current"]), fmt_qty(r["predicted"]), fmt_qty(r["remaining"]), fmt_qty(r["min"]), fmt_qty(r["recommended"]), r["status"]) for r in last["rows"]]
    return make_xlsx("hasil_prediksi_prophet.xlsx", ["Bahan","Stok Saat Ini","Prediksi Pemakaian (Horizon)","Perkiraan Sisa","Minimum","Rekomendasi Restock","Status"], data, "Hasil Prediksi Prophet", f"Data latihan: {last['start']} s.d. {last['end']} | Prediksi: {last['horizon']} hari ke depan | Metode: {last['method']}")


@app.route("/export/prediksi.pdf")
@auth_required
def export_prediksi_pdf():
    last = session.get("last_forecast")
    if not last: flash("Jalankan prediksi terlebih dahulu."); return redirect(url_for("prediksi"))
    data = [(r["material"], f"{fmt_qty(r['current'])} {r['unit']}", f"{fmt_qty(r['predicted'])} {r['unit']}", f"{fmt_qty(r['remaining'])} {r['unit']}", f"{fmt_qty(r['min'])} {r['unit']}", f"{fmt_qty(r['recommended'])} {r['unit']}", r["status"]) for r in last["rows"]]
    return make_pdf("hasil_prediksi_prophet.pdf", "Hasil Prediksi Prophet", ["Bahan","Stok Saat Ini","Prediksi Pemakaian","Perkiraan Sisa","Minimum","Rekomendasi","Status"], data, f"Data latihan: {last['start']} s.d. {last['end']} | Prediksi: {last['horizon']} hari ke depan | Metode: {last['method']}")


init_db()
if __name__ == "__main__":
    app.run(debug=True)
