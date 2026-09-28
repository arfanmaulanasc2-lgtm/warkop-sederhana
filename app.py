from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
import sqlite3, os, io, math, re
from datetime import datetime, date
from functools import wraps

app = Flask(__name__)
app.secret_key = "warkop-sederhana-secret-2026"
BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(BASE, "warkop.db")

MENUS = [
    ("Kopi Hitam", 5000, "gelas", "Kopi bubuk, gula pasir, air mineral", "coffee"),
    ("Nasi Goreng", 15000, "porsi", "Beras, telur, bawang, kecap, minyak, bumbu", "nasi"),
    ("Teh Manis", 4000, "gelas", "Teh celup, gula pasir, air mineral", "tea"),
    ("Good Day", 5000, "gelas", "Good Day sachet, air mineral", "goodday"),
    ("Pisang Goreng", 8000, "porsi", "Pisang, tepung, gula, minyak", "banana"),
    ("Mie Goreng", 10000, "porsi", "Mie instan, telur, minyak, bumbu mie", "noodle"),
    ("Es Jeruk", 6000, "gelas", "Jeruk, gula pasir, air mineral", "orange"),
    ("Air Mineral", 4000, "botol", "Air mineral botol", "water"),
]

RECIPES = {
    "Kopi Hitam": [("Kopi bubuk", "gram", 12), ("Gula pasir", "gram", 10), ("Air mineral", "ml", 180)],
    "Nasi Goreng": [("Beras", "gram", 100), ("Telur", "butir", 1), ("Bawang putih", "gram", 5), ("Bawang merah", "gram", 8), ("Kecap manis", "ml", 10), ("Minyak goreng", "ml", 15), ("Bumbu nasi goreng", "gram", 10)],
    "Teh Manis": [("Teh celup", "sachet", 1), ("Gula pasir", "gram", 15), ("Air mineral", "ml", 200)],
    "Good Day": [("Good Day sachet", "sachet", 1), ("Air mineral", "ml", 180)],
    "Pisang Goreng": [("Pisang", "buah", 3), ("Tepung terigu", "gram", 50), ("Gula pasir", "gram", 5), ("Minyak goreng", "ml", 40)],
    "Mie Goreng": [("Mie instan", "bungkus", 1), ("Telur", "butir", 1), ("Minyak goreng", "ml", 10), ("Bumbu mie", "gram", 5)],
    "Es Jeruk": [("Jeruk peras", "buah", 1), ("Gula pasir", "gram", 15), ("Air mineral", "ml", 180)],
    "Air Mineral": [("Air mineral botol", "botol", 1)],
}

STOCK = [
    ("Kopi bubuk", "gram", 1200, 500), ("Gula pasir", "gram", 3500, 700), ("Air mineral", "ml", 30000, 5000),
    ("Beras", "gram", 9000, 2500), ("Telur", "butir", 35, 10), ("Bawang putih", "gram", 900, 150),
    ("Bawang merah", "gram", 1200, 200), ("Kecap manis", "ml", 3000, 500), ("Minyak goreng", "ml", 1200, 1500),
    ("Bumbu nasi goreng", "gram", 500, 150), ("Teh celup", "sachet", 80, 20), ("Good Day sachet", "sachet", 18, 20),
    ("Pisang", "buah", 18, 15), ("Tepung terigu", "gram", 4000, 700), ("Mie instan", "bungkus", 40, 10),
    ("Bumbu mie", "gram", 300, 50), ("Jeruk peras", "buah", 28, 8), ("Air mineral botol", "botol", 35, 8)
]


def connect():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c


def ensure_column(c, table, column, definition):
    cols = [r[1] for r in c.execute(f"PRAGMA table_info({table})").fetchall()]
    if column not in cols:
        c.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_db():
    c = connect()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE, password TEXT);
    CREATE TABLE IF NOT EXISTS menus(id INTEGER PRIMARY KEY, name TEXT UNIQUE, price INTEGER, unit TEXT, description TEXT, image TEXT);
    CREATE TABLE IF NOT EXISTS stock(id INTEGER PRIMARY KEY, name TEXT UNIQUE, unit TEXT, qty REAL, min_qty REAL);
    CREATE TABLE IF NOT EXISTS recipes(id INTEGER PRIMARY KEY, menu_name TEXT, material TEXT, unit TEXT, qty REAL);
    CREATE TABLE IF NOT EXISTS sales(id INTEGER PRIMARY KEY, invoice TEXT, sale_date TEXT, total INTEGER);
    CREATE TABLE IF NOT EXISTS sale_items(id INTEGER PRIMARY KEY, sale_id INTEGER, menu_name TEXT, qty INTEGER, price INTEGER);
    CREATE TABLE IF NOT EXISTS stock_logs(id INTEGER PRIMARY KEY, material TEXT, type TEXT, qty REAL, unit TEXT, note TEXT, log_date TEXT, menu_name TEXT);
    CREATE TABLE IF NOT EXISTS forecast_runs(id INTEGER PRIMARY KEY, start_date TEXT, end_date TEXT, horizon INTEGER, method TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS forecast_items(id INTEGER PRIMARY KEY, run_id INTEGER, material TEXT, unit TEXT, current REAL, used REAL DEFAULT 0, min_qty REAL, predicted REAL, recommended REAL, status TEXT);
    """)
    # Migrate databases created by older versions.
    ensure_column(c, "stock_logs", "unit", "TEXT")
    ensure_column(c, "stock_logs", "menu_name", "TEXT")
    ensure_column(c, "forecast_items", "used", "REAL DEFAULT 0")
    # Fill missing units in old stock logs from current stock table.
    c.execute("UPDATE stock_logs SET unit=(SELECT unit FROM stock WHERE stock.name=stock_logs.material) WHERE unit IS NULL OR unit=''" )
    c.execute("INSERT OR IGNORE INTO users(username,password) VALUES('admin','admin123')")
    for name, price, unit, desc, img in MENUS:
        c.execute("INSERT OR IGNORE INTO menus(name,price,unit,description,image) VALUES(?,?,?,?,?)", (name, price, unit, desc, img))
    for name, unit, qty, min_qty in STOCK:
        c.execute("INSERT OR IGNORE INTO stock(name,unit,qty,min_qty) VALUES(?,?,?,?)", (name, unit, qty, min_qty))
    for menu, items in RECIPES.items():
        for material, unit, qty in items:
            row = c.execute("SELECT id FROM recipes WHERE menu_name=? AND material=?", (menu, material)).fetchone()
            if not row:
                c.execute("INSERT INTO recipes(menu_name,material,unit,qty) VALUES(?,?,?,?)", (menu, material, unit, qty))
    # Backfill old log menu names from the note where possible.
    old_logs = c.execute("SELECT id,note FROM stock_logs WHERE (menu_name IS NULL OR menu_name='') AND note LIKE 'Penjualan %'").fetchall()
    for r in old_logs:
        m = re.match(r"Penjualan (.+?) x[0-9.]+$", r["note"] or "")
        if m:
            c.execute("UPDATE stock_logs SET menu_name=? WHERE id=?", (m.group(1), r["id"]))
    c.commit()
    c.close()


def auth_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("user"):
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapper


def parse_date_range(default_month=True):
    today = date.today()
    default_start = today.replace(day=1).isoformat() if default_month else ""
    default_end = today.isoformat() if default_month else ""
    start = request.args.get("start") or default_start
    end = request.args.get("end") or default_end
    return start, end


def validate_range(start, end):
    try:
        d1 = datetime.strptime(start, "%Y-%m-%d").date()
        d2 = datetime.strptime(end, "%Y-%m-%d").date()
        return d1, d2
    except ValueError:
        return None, None


def date_range_days(start, end):
    d1, d2 = validate_range(start, end)
    return max(1, (d2 - d1).days + 1) if d1 and d2 else 1


def money(n):
    return "Rp {:,.0f}".format(n).replace(",", ".")


def fmt_qty(n):
    n = float(n or 0)
    return str(int(n)) if n.is_integer() else f"{n:.1f}"


def get_sales_rows(start, end):
    c = connect()
    rows = c.execute("""SELECT s.invoice,s.sale_date,si.menu_name,si.qty,si.price,(si.qty*si.price) line_total
                        FROM sale_items si JOIN sales s ON s.id=si.sale_id
                        WHERE date(s.sale_date) BETWEEN date(?) AND date(?)
                        ORDER BY s.sale_date DESC,s.id DESC,si.id""", (start, end)).fetchall()
    c.close()
    return rows


def get_bahan_detail(start, end):
    c = connect()
    rows = c.execute("""SELECT date(log_date) log_day, COALESCE(menu_name,'-') menu_name, material,
                               COALESCE(unit,(SELECT unit FROM stock s WHERE s.name=stock_logs.material),'') unit,
                               SUM(qty) qty
                        FROM stock_logs
                        WHERE type='KELUAR' AND date(log_date) BETWEEN date(?) AND date(?)
                        GROUP BY date(log_date), menu_name, material, unit
                        ORDER BY log_day DESC, menu_name, material""", (start, end)).fetchall()
    c.close()
    return rows


def get_bahan_rekap_menu(start, end):
    c = connect()
    rows = c.execute("""SELECT COALESCE(menu_name,'-') menu_name, material,
                               COALESCE(unit,(SELECT unit FROM stock s WHERE s.name=stock_logs.material),'') unit,
                               SUM(qty) qty
                        FROM stock_logs
                        WHERE type='KELUAR' AND date(log_date) BETWEEN date(?) AND date(?)
                        GROUP BY menu_name, material, unit
                        ORDER BY menu_name, material""", (start, end)).fetchall()
    c.close()
    return rows


def get_bahan_rekap_material(start, end):
    c = connect()
    rows = c.execute("""SELECT material,
                               COALESCE(unit,(SELECT unit FROM stock s WHERE s.name=stock_logs.material),'') unit,
                               SUM(qty) qty
                        FROM stock_logs
                        WHERE type='KELUAR' AND date(log_date) BETWEEN date(?) AND date(?)
                        GROUP BY material, unit
                        ORDER BY material""", (start, end)).fetchall()
    c.close()
    return rows


def get_grouped_recipes():
    c = connect()
    rows = c.execute("""SELECT r.menu_name,r.material,r.unit,r.qty,m.price,m.unit menu_unit
                        FROM recipes r LEFT JOIN menus m ON m.name=r.menu_name
                        ORDER BY m.id,r.id""").fetchall()
    c.close()
    groups = []
    by_menu = {}
    for r in rows:
        if r["menu_name"] not in by_menu:
            group = {"menu_name": r["menu_name"], "menu_unit": r["menu_unit"], "price": r["price"], "items": []}
            by_menu[r["menu_name"]] = group
            groups.append(group)
        by_menu[r["menu_name"]]["items"].append(r)
    return groups


def daily_menu_sales(start, end):
    c = connect()
    rows = c.execute("""SELECT date(s.sale_date) sale_day, si.menu_name, SUM(si.qty) qty
                      FROM sale_items si JOIN sales s ON s.id=si.sale_id
                      WHERE date(s.sale_date) BETWEEN date(?) AND date(?)
                      GROUP BY date(s.sale_date),si.menu_name ORDER BY sale_day""", (start, end)).fetchall()
    c.close()
    return rows


def opening_stocks(start):
    """Return stock quantity at the start of a selected period.

    The project keeps the original STOCK values as the simulation baseline and
    every subsequent stock movement is recorded in stock_logs. Reconstructing
    the opening balance from those movements keeps historical reports stable
    even after later sales/restocks change the current stock.
    """
    c = connect()
    baseline = {name: float(qty) for name, unit, qty, min_qty in STOCK}
    rows = c.execute("""SELECT material, COALESCE(SUM(CASE WHEN type='MASUK' THEN qty ELSE -qty END),0) net
                        FROM stock_logs WHERE date(log_date) < date(?) GROUP BY material""", (start,)).fetchall()
    c.close()
    result = dict(baseline)
    for r in rows:
        result[r["material"]] = result.get(r["material"], 0.0) + float(r["net"] or 0)
    return result


def forecast_materials(start, end, horizon):
    d1, d2 = validate_range(start, end)
    if not d1 or not d2 or d2 < d1:
        raise ValueError("Periode prediksi tidak valid.")
    c = connect()
    stocks = c.execute("SELECT * FROM stock ORDER BY name").fetchall()
    all_menu_names = [r["name"] for r in c.execute("SELECT name FROM menus ORDER BY id").fetchall()]
    menu_rows = c.execute("""SELECT si.menu_name,SUM(si.qty) total
                             FROM sale_items si JOIN sales s ON s.id=si.sale_id
                             WHERE date(s.sale_date) BETWEEN date(?) AND date(?)
                             GROUP BY si.menu_name""", (start, end)).fetchall()
    used_rows = c.execute("""SELECT material, COALESCE(unit,'') unit, SUM(qty) used
                            FROM stock_logs
                            WHERE type='KELUAR' AND date(log_date) BETWEEN date(?) AND date(?)
                            GROUP BY material, unit""", (start, end)).fetchall()
    c.close()
    opening = opening_stocks(start)
    days = max(1, (d2 - d1).days + 1)
    menu_totals = {r["menu_name"]: float(r["total"]) for r in menu_rows}
    used_by_material = {r["material"]: float(r["used"]) for r in used_rows}
    daily = daily_menu_sales(start, end)
    menu_pred = {}
    used_prophet = False
    try:
        from prophet import Prophet
        import pandas as pd
        idx = pd.date_range(start=start, end=end, freq="D")
        for menu in all_menu_names:
            values = {str(r["sale_day"]): float(r["qty"]) for r in daily if r["menu_name"] == menu}
            df = pd.DataFrame({"ds": idx, "y": [values.get(d.strftime("%Y-%m-%d"), 0.0) for d in idx]})
            if df["y"].sum() <= 0 or len(df) < 2:
                menu_pred[menu] = menu_totals.get(menu, 0.0) / days * horizon
                continue
            model = Prophet(daily_seasonality=False, weekly_seasonality=True, yearly_seasonality=False)
            model.fit(df)
            future = model.make_future_dataframe(periods=horizon, include_history=False)
            pred = model.predict(future)["yhat"].clip(lower=0).sum()
            menu_pred[menu] = float(pred)
            used_prophet = True
    except Exception:
        for menu in all_menu_names:
            menu_pred[menu] = menu_totals.get(menu, 0.0) / days * horizon
    output = []
    for st in stocks:
        used = round(used_by_material.get(st["name"], 0.0), 1)
        start_stock = round(opening.get(st["name"], float(st["qty"])), 1)
        before_restock = round(start_stock - used, 1)
        predicted = 0.0
        for menu, pred in menu_pred.items():
            for material, unit, qty in RECIPES.get(menu, []):
                if material == st["name"]:
                    predicted += pred * qty
        predicted = round(predicted, 1)
        recommended = max(0, math.ceil(predicted + float(st["min_qty"]) - before_restock))
        remaining = round(before_restock - predicted, 1)
        output.append({
            "material": st["name"], "unit": st["unit"], "current": float(st["qty"]),
            "start_stock": start_stock, "used": used, "before_restock": before_restock,
            "min": float(st["min_qty"]), "predicted": predicted, "remaining": remaining,
            "recommended": recommended, "added": 0.0, "after_restock": before_restock,
            "status": "PERLU RESTOCK" if recommended > 0 else "AMAN"
        })
    return output, used_prophet

def save_forecast(start, end, horizon, method, rows):
    c = connect()
    run_id = c.execute("INSERT INTO forecast_runs(start_date,end_date,horizon,method,created_at) VALUES(?,?,?,?,?)",
                       (start, end, horizon, method, datetime.now().isoformat(timespec="seconds"))).lastrowid
    for r in rows:
        c.execute("""INSERT INTO forecast_items(run_id,material,unit,current,used,min_qty,predicted,recommended,status)
                     VALUES(?,?,?,?,?,?,?,?,?)""", (run_id, r["material"], r["unit"], r["current"], r.get("used",0), r["min"], r["predicted"], r["recommended"], r["status"]))
    c.commit(); c.close()
    return run_id


@app.route("/", methods=["GET", "POST"])
def login():
    if session.get("user"):
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        c = connect()
        row = c.execute("SELECT * FROM users WHERE username=? AND password=?", (request.form["username"], request.form["password"])).fetchone()
        c.close()
        if row:
            session["user"] = row["username"]
            return redirect(url_for("dashboard"))
        flash("Username atau password salah.")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear(); return redirect(url_for("login"))


@app.route("/dashboard")
@auth_required
def dashboard():
    c = connect(); today = date.today(); default_start=today.replace(day=1).isoformat(); default_end=today.isoformat()
    start_date=request.args.get("start_date") or default_start; end_date=request.args.get("end_date") or default_end
    if not validate_range(start_date,end_date)[0] or start_date>end_date:
        start_date,end_date=default_start,default_end
    revenue=c.execute("SELECT COALESCE(SUM(total),0) x FROM sales WHERE date(sale_date) BETWEEN date(?) AND date(?)",(start_date,end_date)).fetchone()["x"]
    tx=c.execute("SELECT COUNT(*) x FROM sales WHERE date(sale_date) BETWEEN date(?) AND date(?)",(start_date,end_date)).fetchone()["x"]
    menu_qty=c.execute("SELECT COALESCE(SUM(si.qty),0) x FROM sale_items si JOIN sales s ON s.id=si.sale_id WHERE date(s.sale_date) BETWEEN date(?) AND date(?)",(start_date,end_date)).fetchone()["x"]
    used_total=c.execute("SELECT COALESCE(SUM(qty),0) x FROM stock_logs WHERE type='KELUAR' AND date(log_date) BETWEEN date(?) AND date(?)",(start_date,end_date)).fetchone()["x"]
    low=c.execute("SELECT COUNT(*) x FROM stock WHERE qty<=min_qty").fetchone()["x"]; safe=c.execute("SELECT COUNT(*) x FROM stock WHERE qty>min_qty").fetchone()["x"]
    low_rows=c.execute("SELECT * FROM stock WHERE qty<=min_qty ORDER BY qty ASC").fetchall(); menus=c.execute("SELECT * FROM menus ORDER BY id").fetchall()
    restock_count=c.execute("SELECT COUNT(*) x FROM stock_logs WHERE type='MASUK' AND note LIKE 'Restock dari rekomendasi Prophet%' AND date(log_date) BETWEEN date(?) AND date(?)",(start_date,end_date)).fetchone()["x"]
    c.close()
    return render_template("dashboard.html",title="Dashboard",revenue=revenue,tx=tx,menu_qty=menu_qty,used_total=used_total,restock_count=restock_count,low=low,safe=safe,low_rows=low_rows,menus=menus,start_date=start_date,end_date=end_date)


@app.route("/kasir", methods=["GET", "POST"])
@auth_required
def kasir():
    c = connect(); menus = c.execute("SELECT * FROM menus ORDER BY id").fetchall()
    if request.method == "POST":
        cart = []
        for m in menus:
            qty = int(request.form.get("qty_" + str(m["id"]), 0) or 0)
            if qty > 0:
                cart.append((m["name"], qty, m["price"]))
        sale_date = request.form.get("sale_date") or date.today().isoformat()
        if not cart:
            c.close(); flash("Pilih minimal satu menu."); return redirect(url_for("kasir"))
        try:
            datetime.strptime(sale_date, "%Y-%m-%d")
        except ValueError:
            c.close(); flash("Tanggal transaksi tidak valid."); return redirect(url_for("kasir"))
        needed = {}
        for name, qty, price in cart:
            for material, unit, per_item in RECIPES[name]:
                needed[material] = needed.get(material, 0) + qty * per_item
        for material, need in needed.items():
            row = c.execute("SELECT qty,unit FROM stock WHERE name=?", (material,)).fetchone()
            if not row or row["qty"] < need:
                c.close(); flash(f"Stok {material} tidak cukup. Tersedia {fmt_qty(row['qty']) if row else '0'} {row['unit'] if row else ''}, dibutuhkan {fmt_qty(need)}."); return redirect(url_for("kasir"))
        total = sum(q * p for _, q, p in cart)
        # Satu checkout = satu invoice. Gunakan nomor pendek dan berurutan.
        existing = c.execute("SELECT invoice FROM sales WHERE invoice GLOB '[0-9][0-9][0-9]*'").fetchall()
        seq = 0
        for er in existing:
            inv = str(er["invoice"] or "")
            if inv.isdigit():
                try:
                    seq = max(seq, int(inv))
                except ValueError:
                    pass
        invoice = f"{seq + 1:03d}"
        sid = c.execute("INSERT INTO sales(invoice,sale_date,total) VALUES(?,?,?)", (invoice, sale_date, total)).lastrowid
        for name, qty, price in cart:
            c.execute("INSERT INTO sale_items(sale_id,menu_name,qty,price) VALUES(?,?,?,?)", (sid, name, qty, price))
            for material, unit, per_item in RECIPES[name]:
                used = qty * per_item
                c.execute("UPDATE stock SET qty=qty-? WHERE name=?", (used, material))
                c.execute("INSERT INTO stock_logs(material,type,qty,unit,note,log_date,menu_name) VALUES(?,?,?,?,?,?,?)",
                          (material, "KELUAR", used, unit, f"Penjualan {name} x{qty}", sale_date + " 12:00:00", name))
        c.commit(); c.close()
        flash(f"Transaksi {invoice} tersimpan. Grand Total {money(total)}.")
        return redirect(url_for("kasir"))
    c.close()
    return render_template("kasir.html", title="Kasir", menus=menus, today=date.today().isoformat())


@app.route("/master-menu")
@auth_required
def master_menu():
    c = connect(); menus = c.execute("SELECT * FROM menus ORDER BY id").fetchall(); c.close()
    return render_template("master_menu.html", title="Master Menu", menus=menus)


def get_grouped_sales(start, end):
    rows = get_sales_rows(start, end)
    groups = []
    by_key = {}
    for r in rows:
        day = str(r["sale_date"])[:10]
        key = (day, r["invoice"])
        if key not in by_key:
            group = {"day": day, "invoice": r["invoice"], "items": [], "total_qty": 0, "total": 0}
            by_key[key] = group
            groups.append(group)
        item = {
            "menu_name": r["menu_name"], "qty": int(r["qty"]),
            "price": float(r["price"]), "line_total": float(r["line_total"])
        }
        by_key[key]["items"].append(item)
        by_key[key]["total_qty"] += int(r["qty"])
        by_key[key]["total"] += float(r["line_total"])

    # Tanggal ASC untuk laporan, invoice mengikuti urutan transaksi.
    groups.sort(key=lambda g: (g["day"], int(g["invoice"]) if str(g["invoice"]).isdigit() else str(g["invoice"])))
    return groups

def sales_summary(groups):
    total_tx = len(groups)
    total_qty = sum(g["total_qty"] for g in groups)
    total_sales = sum(g["total"] for g in groups)
    menu_counts = {}
    for g in groups:
        for item in g["items"]:
            menu_counts[item["menu_name"]] = menu_counts.get(item["menu_name"], 0) + item["qty"]
    return total_tx, total_qty, total_sales, menu_counts

@app.route("/penjualan")
@auth_required
def penjualan():
    start,end=parse_date_range(True)
    groups=get_grouped_sales(start,end)
    total_tx,total_qty,total_sales,menu_counts=sales_summary(groups)
    return render_template("penjualan.html",title="Riwayat Penjualan",groups=groups,start=start,end=end,total_tx=total_tx,total_qty=total_qty,total_sales=total_sales,total_menu=sum(len(g["items"]) for g in groups),menu_counts=menu_counts)

@app.route("/stok")
@auth_required
def stok():
    start,end=parse_date_range(True); c=connect(); rows=c.execute("SELECT * FROM stock ORDER BY name").fetchall(); movements=c.execute("""SELECT date(log_date) log_day, material, type, qty, unit, note, COALESCE(menu_name,'-') menu_name FROM stock_logs WHERE date(log_date) BETWEEN date(?) AND date(?) ORDER BY log_day DESC,id DESC""",(start,end)).fetchall(); c.close()
    return render_template("stok.html",title="Stok Bahan",rows=rows,movements=movements,start=start,end=end,today=date.today().isoformat())

@app.route("/stok/tambah",methods=["POST"])
@auth_required
def tambah_stok():
    material=request.form["material"]; qty=float(request.form["qty"]); note=request.form.get("note") or "Stok masuk"; log_date=request.form.get("log_date") or date.today().isoformat(); c=connect(); row=c.execute("SELECT unit FROM stock WHERE name=?",(material,)).fetchone(); unit=row["unit"] if row else ""; c.execute("UPDATE stock SET qty=qty+? WHERE name=?",(qty,material)); c.execute("INSERT INTO stock_logs(material,type,qty,unit,note,log_date,menu_name) VALUES(?,?,?,?,?,?,NULL)",(material,"MASUK",qty,unit,note,log_date+" 12:00:00")); c.commit(); c.close(); flash("Stok berhasil ditambahkan."); return redirect(url_for("stok"))

@app.route("/bahan-terpakai")
@auth_required
def bahan_terpakai():
    start,end=parse_date_range(True); menu_rows=get_bahan_rekap_menu(start,end); material_rows=get_bahan_rekap_material(start,end); detail_rows=get_bahan_detail(start,end); total_used=sum(float(r["qty"]) for r in material_rows)
    return render_template("bahan.html",title="Bahan Terpakai",start=start,end=end,menu_rows=menu_rows,material_rows=material_rows,detail_rows=detail_rows,total_used=total_used)


@app.route("/resep")
@auth_required
def resep():
    groups = get_grouped_recipes()
    return render_template("resep.html", title="Resep Menu", groups=groups)


@app.route("/prediksi", methods=["GET", "POST"])
@auth_required
def prediksi():
    result = []; start = ""; end = ""; horizon = 7; method = ""; run_id = None
    if request.method == "POST":
        start = request.form["start"]; end = request.form["end"]; horizon = int(request.form["horizon"])
        d1, d2 = validate_range(start, end)
        if not d1 or not d2 or d2 < d1:
            flash("Periode tanggal prediksi tidak valid."); return redirect(url_for("prediksi"))
        result, used = forecast_materials(start, end, horizon)
        method = "Prophet" if used else "Rata-rata historis"
        run_id = save_forecast(start, end, horizon, method, result)
        session["last_forecast"] = {"start": start, "end": end, "horizon": horizon, "rows": result, "method": method, "run_id": run_id}
    else:
        last = session.get("last_forecast")
        if last:
            start = last.get("start", "")
            end = last.get("end", "")
            horizon = int(last.get("horizon", 7) or 7)
            method = last.get("method", "")
            run_id = last.get("run_id")
            result = last.get("rows") or []

            # Backward compatibility: older session data may contain forecast
            # rows without the new stock-calculation fields (e.g. start_stock).
            # Rebuild the result from the existing database/current forecast
            # inputs instead of letting Jinja fail on a missing attribute.
            required_keys = {"material", "unit", "start_stock", "used",
                             "before_restock", "predicted", "min",
                             "recommended", "added", "after_restock", "status"}
            if not result or any(not isinstance(r, dict) or not required_keys.issubset(r.keys()) for r in result):
                try:
                    result, used_prophet = forecast_materials(start, end, horizon)
                    if not method:
                        method = "Prophet" if used_prophet else "Rata-rata historis"
                    session["last_forecast"] = {
                        "start": start, "end": end, "horizon": horizon,
                        "rows": result, "method": method, "run_id": run_id
                    }
                    session.modified = True
                except Exception:
                    # If the old session does not contain a usable period,
                    # fall back to the normal default date range.
                    start, end = parse_date_range(True)
                    result = []
                    horizon = 7
                    method = ""
                    run_id = None
        else:
            start, end = parse_date_range(True)
    c = connect(); history = c.execute("SELECT * FROM forecast_runs ORDER BY id DESC LIMIT 20").fetchall(); c.close()
    return render_template("prediksi.html", title="Prediksi Prophet", result=result, start=start, end=end, horizon=horizon, method=method, history=history, run_id=run_id)


@app.route("/prediksi/terapkan", methods=["POST"])
@auth_required
def terapkan():
    last=session.get("last_forecast"); restock_date=request.form.get("restock_date") or date.today().isoformat()
    c=connect()
    if not last:
        c.close(); flash("Jalankan Prophet terlebih dahulu."); return redirect(url_for("prediksi"))
    for key,value in request.form.items():
        if key.startswith("rec_"):
            qty=float(value or 0); material=key[4:]
            if qty>0:
                row=c.execute("SELECT unit,qty FROM stock WHERE name=?",(material,)).fetchone(); unit=row["unit"] if row else ""; before=row["qty"] if row else 0
                c.execute("UPDATE stock SET qty=qty+? WHERE name=?",(qty,material))
                c.execute("INSERT INTO stock_logs(material,type,qty,unit,note,log_date,menu_name) VALUES(?,?,?,?,?,?,NULL)",(material,"MASUK",qty,unit,"Restock dari rekomendasi Prophet",restock_date+" 12:00:00"))
    c.commit(); c.close(); flash("Rekomendasi yang diisi sudah masuk ke stok dan tercatat sebagai restock Prophet."); return redirect(url_for("stok"))


@app.route("/reset-data", methods=["GET", "POST"])
@auth_required
def reset_data():
    if request.method == "GET":
        return render_template("reset.html", title="Reset Data Simulasi")
    if request.form.get("confirmation", "").strip().upper() != "RESET":
        flash("Reset dibatalkan. Ketik RESET untuk mengonfirmasi.")
        return redirect(url_for("dashboard"))
    c = connect()
    c.execute("DELETE FROM sale_items")
    c.execute("DELETE FROM sales")
    c.execute("DELETE FROM stock_logs")
    c.execute("DELETE FROM forecast_items")
    c.execute("DELETE FROM forecast_runs")
    for name, unit, qty, min_qty in STOCK:
        c.execute("UPDATE stock SET qty=?, unit=?, min_qty=? WHERE name=?", (qty, unit, min_qty, name))
    c.commit()
    c.close()
    session.pop("last_forecast", None)
    flash("Data simulasi berhasil direset. Riwayat kembali 0 dan stok dikembalikan ke stok awal.")
    return redirect(url_for("dashboard"))


# ---------- Exports ----------
def make_xlsx(filename, headers, rows, title="Laporan", subtitle="", summary=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    wb=Workbook(); ws=wb.active; ws.title=title[:31]
    ws.append([title]); ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers)); ws["A1"].font=Font(bold=True,size=14); ws["A1"].alignment=Alignment(horizontal="left")
    rownum=2
    if subtitle: ws.append([subtitle]); ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers)); ws["A2"].alignment=Alignment(horizontal="left"); rownum=3
    if summary:
        for label,value in summary: ws.append([label,value]); rownum+=1
    header_row=rownum; ws.append(headers); fill=PatternFill("solid",fgColor="754622"); white=Font(color="FFFFFF",bold=True)
    for cell in ws[header_row]: cell.fill=fill; cell.font=white; cell.alignment=Alignment(horizontal="center")
    for row in rows: ws.append(list(row))
    thin=Side(style="thin",color="DDDDDD")
    for r in ws.iter_rows(min_row=header_row,max_row=ws.max_row):
        for cell in r: cell.border=Border(bottom=thin); cell.alignment=Alignment(vertical="top")
    for col in range(1,len(headers)+1):
        letter=get_column_letter(col); max_len=max([len(str(ws.cell(r,col).value or "")) for r in range(1,ws.max_row+1)]+[10]); ws.column_dimensions[letter].width=min(max_len+3,35)
    bio=io.BytesIO(); wb.save(bio); bio.seek(0); return send_file(bio,as_attachment=True,download_name=filename,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def make_pdf(filename,title,headers,rows,subtitle="",summary=None):
    """Create a landscape administrative PDF with full-width tables and continuation labels."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.enums import TA_LEFT,TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer
    from reportlab.pdfbase.pdfmetrics import stringWidth

    page_w,page_h=landscape(A4)
    left_margin=18
    right_margin=18
    top_margin=38
    bottom_margin=20
    usable_w=page_w-left_margin-right_margin

    bio=io.BytesIO()

    # Keep a little extra room at the top so continuation pages can show
    # "LAMPIRAN 2", "LAMPIRAN 3", etc. without touching the table.
    doc=SimpleDocTemplate(
        bio,
        pagesize=landscape(A4),
        rightMargin=right_margin,
        leftMargin=left_margin,
        topMargin=top_margin,
        bottomMargin=bottom_margin,
        title=title,
        author="Warkop Sederhana"
    )

    styles=getSampleStyleSheet()
    title_style=ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        alignment=TA_LEFT,
        spaceAfter=4,
    )
    sub_style=ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        alignment=TA_LEFT,
        spaceAfter=7,
    )
    cell_style=ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9,
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    header_style=ParagraphStyle(
        "TableHeader",
        parent=cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
        fontSize=7.3,
        leading=8.5,
        alignment=TA_LEFT,
    )
    summary_label_style=ParagraphStyle(
        "SummaryLabel",
        parent=cell_style,
        fontName="Helvetica-Bold",
    )

    story=[Paragraph(str(title),title_style)]
    if subtitle:
        story.append(Paragraph(str(subtitle),sub_style))

    # Summary uses the full report width instead of a small centered block.
    if summary:
        summary_data=[]
        for a,b in summary:
            summary_data.append([
                Paragraph(str(a),summary_label_style),
                Paragraph(str(b),cell_style)
            ])
        summary_table=Table(
            summary_data,
            colWidths=[usable_w*0.28,usable_w*0.72],
            hAlign="LEFT",
            repeatRows=0,
        )
        summary_table.setStyle(TableStyle([
            ("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#D5D5D5")),
            ("BACKGROUND",(0,0),(0,-1),colors.HexColor("#F5F5F5")),
            ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
            ("LEFTPADDING",(0,0),(-1,-1),5),
            ("RIGHTPADDING",(0,0),(-1,-1),5),
            ("TOPPADDING",(0,0),(-1,-1),4),
            ("BOTTOMPADDING",(0,0),(-1,-1),4),
        ]))
        story.extend([summary_table,Spacer(1,8)])

    # Convert every cell to Paragraphs so long values wrap instead of
    # forcing the table outside the page.
    raw_headers=[str(h) for h in headers]
    raw_rows=[[str(v) for v in row] for row in rows]

    # Estimate a sensible width for each column from headers and sample data.
    # Then scale all columns proportionally to fill almost the entire page.
    ncols=len(raw_headers)
    if ncols:
        desired=[]
        for i,h in enumerate(raw_headers):
            samples=[h]
            samples.extend(r[i] for r in raw_rows[:80] if i < len(r))
            longest=max((stringWidth(x,"Helvetica",7.2) for x in samples),default=25)
            header_w=stringWidth(h,"Helvetica-Bold",7.2)+14
            desired.append(max(longest+14,header_w,42))

        # Give naturally long-text columns more room, but never let one
        # column consume the entire report.
        desired=[min(w,usable_w*0.32) for w in desired]
        total=sum(desired)
        if total <= 0:
            col_widths=[usable_w/ncols]*ncols
        else:
            col_widths=[w*usable_w/total for w in desired]
            # Avoid extremely narrow columns after scaling.
            min_w=42
            if any(w < min_w for w in col_widths) and usable_w >= min_w*ncols:
                fixed=[]
                deficit=0
                free=0
                for w in col_widths:
                    if w < min_w:
                        fixed.append(min_w); deficit += min_w-w
                    else:
                        fixed.append(w); free += w
                if free > deficit:
                    factor=(free-deficit)/free
                    col_widths=[min_w if w==min_w else w*factor for w in fixed]
                else:
                    col_widths=[usable_w/ncols]*ncols

        # Correct floating-point rounding so the table ends exactly at the
        # usable right edge rather than leaving a large blank area.
        diff=usable_w-sum(col_widths)
        col_widths[-1]+=diff

        data=[[Paragraph(h,header_style) for h in raw_headers]]
        for row in raw_rows:
            padded=list(row)+[""]*max(0,ncols-len(row))
            data.append([Paragraph(padded[i],cell_style) for i in range(ncols)])

        table=Table(
            data,
            colWidths=col_widths,
            repeatRows=1,
            hAlign="LEFT",
            splitByRow=1,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#7A4724")),
            ("TEXTCOLOR",(0,0),(-1,0),colors.white),
            ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
            ("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#CFCFCF")),
            ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
            ("LEFTPADDING",(0,0),(-1,-1),4),
            ("RIGHTPADDING",(0,0),(-1,-1),4),
            ("TOPPADDING",(0,0),(-1,-1),4),
            ("BOTTOMPADDING",(0,0),(-1,-1),4),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#FAFAFA")]),
        ]))
        story.append(table)

    def draw_continuation(canvas,doc):
        # Page 1 keeps the normal report header. Pages 2+ receive a clear
        # continuation label in the upper-right corner.
        if doc.page > 1:
            canvas.saveState()
            canvas.setFont("Helvetica-Bold",8.5)
            canvas.setFillColor(colors.HexColor("#555555"))
            canvas.drawRightString(page_w-right_margin, page_h-19, f"LAMPIRAN {doc.page}")
            canvas.restoreState()

    doc.build(story,onFirstPage=draw_continuation,onLaterPages=draw_continuation)
    bio.seek(0)
    return send_file(bio,as_attachment=True,download_name=filename,mimetype="application/pdf")

@app.route("/export/penjualan.xlsx")
@auth_required
def export_penjualan_xlsx():
    start, end = parse_date_range(True)
    groups = get_grouped_sales(start, end)
    total_tx,total_qty,total_sales,menu_counts=sales_summary(groups)
    rows=[]
    for g in groups:
        first=True
        for item in g["items"]:
            rows.append((g["invoice"] if first else "", g["day"] if first else "", item["menu_name"], item["qty"], money(item["price"]), money(item["line_total"])))
            first=False
        rows.append(("", "", f"TOTAL INVOICE {g['invoice']}", g["total_qty"], "", money(g["total"])))
    rows.append(("", "", "", "", "", ""))
    rows.append(("", "", "GRAND TOTAL TRANSAKSI", total_tx, "", money(total_sales)))
    rows.append(("", "", "TOTAL ITEM TERJUAL", total_qty, "", ""))
    for menu, qty in sorted(menu_counts.items()):
        rows.append(("", "", menu, qty, "", ""))
    summary=[("Total transaksi/invoice",total_tx),("Total qty/item terjual",total_qty),("Total penjualan",money(total_sales))]
    return make_xlsx("riwayat_penjualan.xlsx",["Invoice","Tanggal","Menu / Keterangan","Qty","Harga Satuan","Total"],rows,"Riwayat Penjualan",f"Periode: {start} s.d. {end}",summary)

@app.route("/export/penjualan.pdf")
@auth_required
def export_penjualan_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether

    start,end=parse_date_range(True)
    groups=get_grouped_sales(start,end)
    total_tx,total_qty,total_sales,menu_counts=sales_summary(groups)
    page_w,page_h=landscape(A4); left=18; right=18; top=40; bottom=20; usable=page_w-left-right
    bio=io.BytesIO()
    doc=SimpleDocTemplate(bio,pagesize=landscape(A4),leftMargin=left,rightMargin=right,topMargin=top,bottomMargin=bottom,title="Riwayat Penjualan",author="Warkop Sederhana")
    styles=getSampleStyleSheet()
    title_style=ParagraphStyle("rt",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=15,leading=18,alignment=TA_LEFT,spaceAfter=4)
    sub_style=ParagraphStyle("rs",parent=styles["Normal"],fontSize=8.5,leading=11,alignment=TA_LEFT,spaceAfter=7)
    date_style=ParagraphStyle("rd",parent=styles["Heading2"],fontName="Helvetica-Bold",fontSize=10.5,leading=13,alignment=TA_LEFT,spaceBefore=7,spaceAfter=4)
    inv_style=ParagraphStyle("ri",parent=styles["Heading3"],fontName="Helvetica-Bold",fontSize=8.8,leading=11,alignment=TA_LEFT,spaceBefore=3,spaceAfter=3)
    cell=ParagraphStyle("rc",parent=styles["Normal"],fontSize=7.5,leading=9,alignment=TA_LEFT)
    cell_r=ParagraphStyle("rcr",parent=cell,alignment=TA_RIGHT)
    head=ParagraphStyle("rh",parent=cell,fontName="Helvetica-Bold",textColor=colors.white)
    story=[Paragraph("Riwayat Penjualan",title_style),Paragraph(f"Periode: {start} s.d. {end}",sub_style)]
    summary_data=[[Paragraph("Total transaksi / invoice",cell),Paragraph(str(total_tx),cell_r)], [Paragraph("Total item / qty terjual",cell),Paragraph(str(total_qty),cell_r)], [Paragraph("Total penjualan",cell),Paragraph(money(total_sales),cell_r)]]
    st=Table(summary_data,colWidths=[usable*0.45,usable*0.55],hAlign="LEFT")
    st.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#D5D5D5")),("BACKGROUND",(0,0),(0,-1),colors.HexColor("#F5F5F5")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
    story.extend([st,Spacer(1,8)])
    current_day=None
    for g in groups:
        if g["day"] != current_day:
            current_day=g["day"]
            story.append(Paragraph(f"Tanggal: {current_day}",date_style))
        story.append(Paragraph(f"Invoice {g['invoice']}",inv_style))
        data=[[Paragraph("Menu",head),Paragraph("Qty",head),Paragraph("Harga Satuan",head),Paragraph("Total",head)]]
        for item in g["items"]:
            data.append([Paragraph(item["menu_name"],cell),Paragraph(str(item["qty"]),cell_r),Paragraph(money(item["price"]),cell_r),Paragraph(money(item["line_total"]),cell_r)])
        data.append([Paragraph(f"Total Invoice {g['invoice']}",ParagraphStyle("tot",parent=cell,fontName="Helvetica-Bold")),Paragraph(str(g["total_qty"]),ParagraphStyle("totr",parent=cell_r,fontName="Helvetica-Bold")),Paragraph("",cell),Paragraph(money(g["total"]),ParagraphStyle("totr2",parent=cell_r,fontName="Helvetica-Bold"))])
        tbl=Table(data,colWidths=[usable*0.52,usable*0.10,usable*0.19,usable*0.19],repeatRows=1,hAlign="LEFT",splitByRow=1)
        tbl.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#7A4724")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#CFCFCF")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BACKGROUND",(0,-1),(-1,-1),colors.HexColor("#F5F5F5")),("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold"),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
        story.extend([tbl,Spacer(1,7)])
    story.append(Paragraph("GRAND TOTAL",date_style))
    grand=[[Paragraph("Total Transaksi / Invoice",cell),Paragraph(str(total_tx),cell_r)], [Paragraph("Total Item / Qty Terjual",cell),Paragraph(str(total_qty),cell_r)], [Paragraph("Total Penjualan",cell),Paragraph(money(total_sales),cell_r)]]
    gt=Table(grand,colWidths=[usable*0.45,usable*0.55],hAlign="LEFT")
    gt.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#BDBDBD")),("BACKGROUND",(0,0),(0,-1),colors.HexColor("#F5F5F5")),("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5)]))
    story.extend([gt,Spacer(1,8),Paragraph("Rekap Menu Terjual",date_style)])
    menu_data=[[Paragraph("Menu",head),Paragraph("Total Terjual",head)]]
    for menu,qty in sorted(menu_counts.items()): menu_data.append([Paragraph(menu,cell),Paragraph(str(qty),cell_r)])
    mt=Table(menu_data,colWidths=[usable*0.72,usable*0.28],repeatRows=1,hAlign="LEFT",splitByRow=1)
    mt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#7A4724")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#CFCFCF")),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
    story.append(mt)
    def draw_page(canvas,doc):
        if doc.page>1:
            canvas.saveState(); canvas.setFont("Helvetica-Bold",8.5); canvas.setFillColor(colors.HexColor("#555555")); canvas.drawRightString(page_w-right,page_h-19,f"LAMPIRAN {doc.page}"); canvas.restoreState()
    doc.build(story,onFirstPage=draw_page,onLaterPages=draw_page)
    bio.seek(0)
    return send_file(bio,as_attachment=True,download_name="riwayat_penjualan.pdf",mimetype="application/pdf")

@app.route("/export/bahan.xlsx")
@auth_required
def export_bahan_xlsx():
    start,end=parse_date_range(True); rows=get_bahan_rekap_material(start,end); data=[(r["material"],r["unit"],fmt_qty(r["qty"])) for r in rows]; return make_xlsx("bahan_terpakai.xlsx",["Bahan","Satuan","Total Terpakai"],data,"Bahan Terpakai",f"Periode: {start} s.d. {end}",[("Total jenis bahan",len(rows)),("Total pemakaian",fmt_qty(sum(float(r["qty"]) for r in rows)))])
@app.route("/export/bahan.pdf")
@auth_required
def export_bahan_pdf():
    start,end=parse_date_range(True); rows=get_bahan_rekap_material(start,end); data=[(r["material"],r["unit"],fmt_qty(r["qty"])) for r in rows]; return make_pdf("bahan_terpakai.pdf","Bahan Terpakai",["Bahan","Satuan","Total Terpakai"],data,f"Periode: {start} s.d. {end}",[("Total jenis bahan",len(rows)),("Total pemakaian",fmt_qty(sum(float(r["qty"]) for r in rows)))])

@app.route("/export/stok.xlsx")
@auth_required
def export_stok_xlsx():
    c=connect(); rows=c.execute("SELECT * FROM stock ORDER BY name").fetchall(); c.close(); data=[(r["name"],r["unit"],fmt_qty(r["qty"]),fmt_qty(r["min_qty"]),"Aman" if r["qty"]>r["min_qty"] else "Menipis") for r in rows]; return make_xlsx("stok_bahan_terkini.xlsx",["Bahan","Satuan","Stok Saat Ini","Minimum","Status"],data,"Stok Bahan Terkini",f"Posisi stok saat ini: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}")
@app.route("/export/stok.pdf")
@auth_required
def export_stok_pdf():
    c=connect(); rows=c.execute("SELECT * FROM stock ORDER BY name").fetchall(); c.close(); data=[(r["name"],r["unit"],fmt_qty(r["qty"]),fmt_qty(r["min_qty"]),"Aman" if r["qty"]>r["min_qty"] else "Menipis") for r in rows]; return make_pdf("stok_bahan_terkini.pdf","Stok Bahan Terkini",["Bahan","Satuan","Stok Saat Ini","Minimum","Status"],data,f"Posisi stok saat ini: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}")

def last_forecast_rows():
    last=session.get("last_forecast")
    if not last: return None
    return last
@app.route("/export/prediksi.xlsx")
@auth_required
def export_prediksi_xlsx():
    last=last_forecast_rows()
    if not last: flash("Jalankan prediksi terlebih dahulu."); return redirect(url_for("prediksi"))
    data=[(r["material"],r["unit"],fmt_qty(r.get("start_stock",0)),fmt_qty(r.get("used",0)),fmt_qty(r.get("before_restock",0)),fmt_qty(r["predicted"]),fmt_qty(r["min"]),fmt_qty(r["recommended"]),fmt_qty(r.get("added",0)),fmt_qty(r.get("after_restock",r.get("before_restock",0))),r["status"]) for r in last["rows"]]
    return make_xlsx("hasil_prediksi_prophet.xlsx",["Bahan","Satuan","Stok Awal","Stok Terpakai","Stok Sebelum Restock","Prediksi Pemakaian","Minimum","Rekomendasi Restock","Jumlah Ditambahkan","Stok Setelah Restock","Status"],data,"Hasil Prediksi Prophet",f"Histori: {last['start']} s.d. {last['end']} | Prediksi: {last['horizon']} hari ke depan | Metode: {last['method']}")

@app.route("/export/prediksi.pdf")
@auth_required
def export_prediksi_pdf():
    last=last_forecast_rows()
    if not last: flash("Jalankan prediksi terlebih dahulu."); return redirect(url_for("prediksi"))
    data=[(r["material"],f"{fmt_qty(r.get('start_stock',0))} {r['unit']}",f"{fmt_qty(r.get('used',0))} {r['unit']}",f"{fmt_qty(r.get('before_restock',0))} {r['unit']}",f"{fmt_qty(r['predicted'])} {r['unit']}",f"{fmt_qty(r['min'])} {r['unit']}",f"{fmt_qty(r['recommended'])} {r['unit']}",f"{fmt_qty(r.get('added',0))} {r['unit']}",f"{fmt_qty(r.get('after_restock',r.get('before_restock',0)))} {r['unit']}",r["status"]) for r in last["rows"]]
    return make_pdf("hasil_prediksi_prophet.pdf","Hasil Prediksi Prophet",["Bahan","Stok Awal","Terpakai","Sebelum Restock","Prediksi","Minimum","Rekomendasi","Ditambahkan","Setelah Restock","Status"],data,f"Histori: {last['start']} s.d. {last['end']} | Prediksi: {last['horizon']} hari ke depan | Metode: {last['method']}")

@app.route("/laporan-stok")
@auth_required
def laporan_stok():
    start,end=parse_date_range(True); report=build_laporan_stok(start,end)
    c=connect(); run=c.execute("SELECT id FROM forecast_runs WHERE start_date=? AND end_date=? ORDER BY id DESC LIMIT 1",(start,end)).fetchone(); pred={r['material']:{'predicted':float(r['predicted']),'recommended':float(r['recommended']),'min':float(r['min_qty'])} for r in c.execute("SELECT material,predicted,recommended,min_qty FROM forecast_items WHERE run_id=?",(run['id'],)).fetchall()} if run else {}
    current={r['name']:float(r['qty']) for r in c.execute("SELECT name,qty FROM stock").fetchall()}
    logs=c.execute("SELECT id,date(log_date) d,material,qty,unit FROM stock_logs WHERE type='MASUK' AND note LIKE 'Restock dari rekomendasi Prophet%' AND date(log_date) BETWEEN date(?) AND date(?) ORDER BY log_date DESC,id DESC",(start,end)).fetchall()
    restocks=[]
    for lg in logs:
        after=current.get(lg['material'],0)
        later=c.execute("SELECT COALESCE(SUM(CASE WHEN type='MASUK' THEN qty ELSE -qty END),0) x FROM stock_logs WHERE material=? AND id>?",(lg['material'],lg['id'])).fetchone()['x']
        after_event=after-later; before_event=after_event-float(lg['qty']); pp=pred.get(lg['material'],{})
        restocks.append({'d':lg['d'],'material':lg['material'],'unit':lg['unit'],'before':before_event,'predicted':pp.get('predicted',0),'min':pp.get('min',0),'recommended':pp.get('recommended',float(lg['qty'])),'added':float(lg['qty']),'after':after_event})
    c.close()
    return render_template("laporan_stok.html",title="Laporan Perbandingan Stok",start=start,end=end,report=report,restocks=restocks)


@app.route("/export/laporan-stok.xlsx")
@auth_required
def export_laporan_stok_xlsx():
    start,end=parse_date_range(True); report=build_laporan_stok(start,end); data=[(r["material"],r["unit"],fmt_qty(r["start_stock"]),fmt_qty(r["used"]),fmt_qty(r["before_restock"]),fmt_qty(r["predicted"]),fmt_qty(r["min"]),fmt_qty(r["recommended"]),fmt_qty(r["added"]),fmt_qty(r["after_restock"]),r["status"]) for r in report]; return make_xlsx("laporan_perbandingan_stok.xlsx",["Bahan","Satuan","Stok Awal","Terpakai","Stok Sebelum Restock","Prediksi Prophet","Minimum","Rekomendasi Restock","Jumlah Ditambahkan","Stok Setelah Restock","Status"],data,"Laporan Perbandingan Stok",f"Periode: {start} s.d. {end}")
@app.route("/export/laporan-stok.pdf")
@auth_required
def export_laporan_stok_pdf():
    start,end=parse_date_range(True); report=build_laporan_stok(start,end); data=[(r["material"],r["unit"],fmt_qty(r["start_stock"]),fmt_qty(r["used"]),fmt_qty(r["before_restock"]),fmt_qty(r["predicted"]),fmt_qty(r["min"]),fmt_qty(r["recommended"]),fmt_qty(r["added"]),fmt_qty(r["after_restock"]),r["status"]) for r in report]; return make_pdf("laporan_perbandingan_stok.pdf","Laporan Perbandingan Stok",["Bahan","Satuan","Stok Awal","Terpakai","Sebelum Restock","Prediksi","Minimum","Rekomendasi","Ditambahkan","Setelah Restock","Status"],data,f"Periode: {start} s.d. {end}")
@app.route("/export/restock-prophet.xlsx")
@auth_required
def export_restock_xlsx():
    start,end=parse_date_range(True); rows=get_restock_report_rows(start,end); data=[(r['d'],r['material'],fmt_qty(r['before']),fmt_qty(r['predicted']),fmt_qty(r['min']),fmt_qty(r['recommended']),fmt_qty(r['added']),fmt_qty(r['after'])) for r in rows]; return make_xlsx("riwayat_restock_prophet.xlsx",["Tanggal","Bahan","Stok Sebelum","Prediksi","Minimum","Rekomendasi","Ditambahkan","Stok Setelah"],data,"Riwayat Restock Prophet",f"Periode: {start} s.d. {end}")
@app.route("/export/restock-prophet.pdf")
@auth_required
def export_restock_pdf():
    start,end=parse_date_range(True); rows=get_restock_report_rows(start,end); data=[(r['d'],r['material'],fmt_qty(r['before']),fmt_qty(r['predicted']),fmt_qty(r['min']),fmt_qty(r['recommended']),fmt_qty(r['added']),fmt_qty(r['after'])) for r in rows]; return make_pdf("riwayat_restock_prophet.pdf","Riwayat Restock Prophet",["Tanggal","Bahan","Stok Sebelum","Prediksi","Minimum","Rekomendasi","Ditambahkan","Stok Setelah"],data,f"Periode: {start} s.d. {end}")

def get_restock_report_rows(start,end):
    c=connect(); run=c.execute("SELECT id FROM forecast_runs WHERE start_date=? AND end_date=? ORDER BY id DESC LIMIT 1",(start,end)).fetchone(); pred={r['material']:{'predicted':float(r['predicted']),'recommended':float(r['recommended']),'min':float(r['min_qty'])} for r in c.execute("SELECT material,predicted,recommended,min_qty FROM forecast_items WHERE run_id=?",(run['id'],)).fetchall()} if run else {}; current={r['name']:float(r['qty']) for r in c.execute("SELECT name,qty FROM stock").fetchall()}; logs=c.execute("SELECT id,date(log_date) d,material,qty,unit FROM stock_logs WHERE type='MASUK' AND note LIKE 'Restock dari rekomendasi Prophet%' AND date(log_date) BETWEEN date(?) AND date(?) ORDER BY log_date DESC,id DESC",(start,end)).fetchall(); out=[]
    for lg in logs:
        after=current.get(lg['material'],0); later=c.execute("SELECT COALESCE(SUM(CASE WHEN type='MASUK' THEN qty ELSE -qty END),0) x FROM stock_logs WHERE material=? AND id>?",(lg['material'],lg['id'])).fetchone()['x']; after_event=after-later; before_event=after_event-float(lg['qty']); pp=pred.get(lg['material'],{}); out.append({'d':lg['d'],'material':lg['material'],'unit':lg['unit'],'before':before_event,'predicted':pp.get('predicted',0),'min':pp.get('min',0),'recommended':pp.get('recommended',float(lg['qty'])),'added':float(lg['qty']),'after':after_event})
    c.close(); return out

def build_laporan_stok(start,end):
    c=connect()
    stocks=c.execute("SELECT * FROM stock ORDER BY name").fetchall()
    used=c.execute("SELECT material,SUM(qty) used FROM stock_logs WHERE type='KELUAR' AND date(log_date) BETWEEN date(?) AND date(?) GROUP BY material",(start,end)).fetchall()
    rest=c.execute("SELECT material,SUM(qty) qty FROM stock_logs WHERE type='MASUK' AND note LIKE 'Restock dari rekomendasi Prophet%' AND date(log_date) BETWEEN date(?) AND date(?) GROUP BY material",(start,end)).fetchall()
    run=c.execute("SELECT id,horizon,method FROM forecast_runs WHERE start_date=? AND end_date=? ORDER BY id DESC LIMIT 1",(start,end)).fetchone()
    pred_rows=c.execute("SELECT material,predicted,recommended FROM forecast_items WHERE run_id=?",(run['id'],)).fetchall() if run else []
    c.close()
    opening=opening_stocks(start)
    um={r['material']:float(r['used']) for r in used}
    rm={r['material']:float(r['qty']) for r in rest}
    pm={r['material']:{'predicted':float(r['predicted']),'recommended':float(r['recommended'])} for r in pred_rows}
    out=[]
    for st in stocks:
        name=st['name']; start_stock=round(opening.get(name,float(st['qty'])),1); u=round(um.get(name,0),1)
        before=round(start_stock-u,1); added=round(rm.get(name,0),1); after=round(before+added,1)
        p=pm.get(name,{}); pred=p.get('predicted',0); rec=p.get('recommended',0)
        out.append({'material':name,'unit':st['unit'],'start_stock':start_stock,'used':u,'before_restock':before,'predicted':pred,'recommended':rec,'added':added,'after_restock':after,'min':float(st['min_qty']),'status':'Aman' if after>float(st['min_qty']) else 'Menipis'})
    return out


init_db()
if __name__ == "__main__":
    app.run(debug=True)
